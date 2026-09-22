"""Генерация презентации проекта (presentation.pptx) через python-pptx.

Запуск (на своей машине):
    pip install python-pptx
    python docs/build_presentation.py
Результат: docs/presentation.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x21, 0x29, 0x5C)
TEAL = RGBColor(0x02, 0x80, 0x90)
MINT = RGBColor(0x02, 0xC3, 0x9A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x27, 0x33)
GREY = RGBColor(0x5A, 0x6B, 0x7B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _box(slide, x, y, w, h):
    return slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame


def _fill(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE

    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp


def _p(tf, text, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, space=6):
    para = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    para.alignment = align; para.space_after = Pt(space)
    run = para.add_run(); run.text = text
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    run.font.name = "Calibri"
    return para


def title_slide(title, subtitle):
    s = prs.slides.add_slide(BLANK)
    _fill(s, 0, 0, 13.333, 7.5, NAVY)
    _fill(s, 0, 3.05, 13.333, 0.04, MINT)
    t = _box(s, 0.9, 2.0, 11.5, 1.4)
    _p(t, title, 40, WHITE, bold=True)
    st = _box(s, 0.9, 3.2, 11.5, 1.2)
    _p(st, subtitle, 20, RGBColor(0xCA, 0xDC, 0xFC))
    return s


def content_slide(title):
    s = prs.slides.add_slide(BLANK)
    h = _box(s, 0.7, 0.45, 12, 1.0)
    _p(h, title, 32, NAVY, bold=True)
    return s


def stat_cards(slide, cards, y=2.0):
    n = len(cards); gap = 0.4
    w = (12.0 - gap * (n - 1)) / n
    for i, (num, label) in enumerate(cards):
        x = 0.7 + i * (w + gap)
        _fill(slide, x, y, w, 2.2, RGBColor(0xEF, 0xF6, 0xF7))
        tf = _box(slide, x + 0.15, y + 0.3, w - 0.3, 1.7)
        _p(tf, num, 40, TEAL, bold=True, align=PP_ALIGN.CENTER)
        _p(tf, label, 13, GREY, align=PP_ALIGN.CENTER)


# 1
title_slide("Категоризация e-commerce товаров",
            "Production ML Spring 2026 · текст + картинки · MaaS · macro F1-score")

# 2
s = content_slide("Задача и метрика")
tf = _box(s, 0.7, 1.6, 11.9, 4.5)
for line in [
    "• Предсказать category_ind товара по контенту: текстовые поля, картинка, URL.",
    "• Категории заданы деревом (tree.csv); у товара ровно одна категория-лист.",
    "• Метрика — macro F1-score: каждый класс важен одинаково.",
    "• Результат — и сабмишн на Kaggle, и работающий сервис /predict.",
]:
    _p(tf, line, 18, DARK, space=12)

# 3
s = content_slide("Данные и EDA")
stat_cards(s, [("111 405", "товаров в train"), ("207", "классов"),
               ("31 563×", "дисбаланс max/min"), ("100%", "есть картинка")])
tf = _box(s, 0.7, 4.5, 11.9, 2.3)
for line in [
    "• Длинный хвост: 7 классов-синглтонов, 30 классов с ≤5 примерами — главный ограничитель macro-F1.",
    "• Поля vendor/url почти всегда заполнены; name пуст у 44%, description — у 65%.",
    "• Вывод: ценны и текст, и картинка → late fusion + работа с дисбалансом.",
]:
    _p(tf, line, 16, DARK, space=8)

# 4
s = content_slide("Идея решения и архитектура модели")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "full_text = vendor + name + model + type_prefix + description + домен + токены URL",
    "1. TF-IDF (word+char) + SGD/SVM — сильный текстовый бейзлайн.",
    "2. E5 (multilingual) эмбеддинги текста + LogReg — семантика.",
    "3. Late fusion: E5 + CLIP (картинки) → MLP-голова.",
    "4. Ансамбль: смешивание вероятностей всех моделей.",
    "Дисбаланс: class_weight=balanced, иерархия дерева, калибровка под macro-F1.",
]:
    _p(tf, line, 18, DARK, space=12)

# 5
s = content_slide("Воспроизводимость и трекинг (MLflow)")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "• Весь эксперимент — одним скриптом: run_end2end.sh (CPU) / train_full_pipeline.sh (GPU).",
    "• Зависимости зафиксированы: uv + pyproject + lock; extra embeddings для DL.",
    "• Отдельные скрипты обработки данных и обучения, конфиги в configs/experiments/*.toml.",
    "• MLflow логирует параметры, метрики (accuracy, macro-F1) и саму модель; артефакты — на S3.",
]:
    _p(tf, line, 18, DARK, space=12)

# 6
s = content_slide("Результаты (Kaggle, macro-F1)")
rows = [("Сабмишн", "Public score"),
        ("первый baseline", "0.288"),
        ("TF-IDF word (run03)", "0.364"),
        ("TF-IDF word+char SGD (run06)", "0.387"),
        ("E5 / E5+CLIP / ансамбль", "ожидается рост")]
table = s.shapes.add_table(len(rows), 2, Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.6)).table
for r, (a, b) in enumerate(rows):
    table.cell(r, 0).text = a; table.cell(r, 1).text = b
    for c in range(2):
        cell = table.cell(r, c)
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            cell.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
            cell.text_frame.paragraphs[0].runs[0].font.bold = True

# 7
s = content_slide("Production-архитектура")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "Один docker compose (./run_service.sh) поднимает весь стек:",
    "• Flask-гейтвей: REST /predict ({url, texts, image_url} → {category_ind}), порт SERVICE_PORT.",
    "• Triton Inference Server: модель как ONNX-ансамбль (E5 → голова).",
    "• Redis: кэш предсказаний.",
    "• Prometheus + Grafana + Alertmanager: метрики, дашборд, алерты.",
    "• Evidently + Airflow: дрифт и регулярное переобучение. Чекпойнты — на S3/MinIO.",
]:
    _p(tf, line, 18, DARK, space=10)

# 8
s = content_slide("MaaS на Triton (ONNX)")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "Ансамбль ecom_ensemble: TEXT → preprocess (токенизация E5) → text_encoder (ONNX) →",
    "category_head (ONNX) → postprocess (argmax) → CATEGORY_IND.",
    "• Конвертер моделей в ONNX: triton/prepare_models.py (паттерн hw7).",
    "• Бэкенд переключается MODEL_BACKEND=triton|local; по умолчанию local (надёжно на CPU без GPU).",
]:
    _p(tf, line, 18, DARK, space=12)

# 9
s = content_slide("Мониторинг и ML-дрифт")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "• Технический мониторинг: Grafana-дашборд (RPS, latency p50/p95, доля 5xx, cache hit).",
    "• Алерты (Alertmanager): сервис недоступен, >5% 5xx, p95 > 1s.",
    "• ML-мониторинг: Evidently сравнивает распределения train vs свежие данные (Data Drift).",
    "• Отчёт о дрифте генерируется регулярно из Airflow.",
]:
    _p(tf, line, 18, DARK, space=12)

# 10
s = content_slide("Автоматизация: Airflow + CI/CD")
tf = _box(s, 0.7, 1.6, 11.9, 5)
for line in [
    "• Airflow DAG ecom_retrain (@daily): extract → features → drift → retrain → publish на S3.",
    "• GitLab CI/CD: lint (ruff) → test (pytest) → build образа сервиса на каждый push в master.",
    "• Тесты: контракт /predict, признаки, иерархия, ансамбль вероятностей.",
]:
    _p(tf, line, 18, DARK, space=12)

# 11
s = content_slide("Карта критериев → реализация")
tf = _box(s, 0.7, 1.5, 11.9, 5.4)
for line in [
    "РК (15): EDA, идея/архитектура, end2end, фикс зависимостей, MLflow, скрипты данных/обучения.",
    "Защита (30): Kaggle-качество, Triton+ONNX, Grafana+алерты, дрифт (Evidently),",
    "Airflow (регулярные расчёты), кэш (Redis), тесты + CI/CD, MaaS, презентация, диаграмма.",
    "Все компоненты опираются на паттерны из ДЗ 4–9 и лекций курса.",
]:
    _p(tf, line, 18, DARK, space=12)

# 12
title_slide("Спасибо!", "Вопросы? · репозиторий: gitlab .../e.puzyreva · README + docs/architecture.svg")

out = Path(__file__).parent / "presentation.pptx"
prs.save(out)
print(f"Saved: {out}")
